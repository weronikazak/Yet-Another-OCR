import os
from flask import Flask, render_template, request, redirect, url_for, send_file
from werkzeug.utils import secure_filename
from logic import find_text
from docx import Document
from fpdf import FPDF

app = Flask(__name__)

app.config['UPLOAD_FOLDER'] = 'static/uploads/'
app.config['ALLOWED_EXTENSIONS'] = set(['png', 'jpg', 'jpeg'])


def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1] in app.config['ALLOWED_EXTENSIONS']


@app.route('/')
def index():
    return render_template('upload.html')


@app.route('/upload', methods=['POST'])
def upload():
    uploaded_files = request.files.getlist("file[]")
    if not any(f for f in uploaded_files):
      return redirect(url_for('index'))
    file_details = []
    for file in uploaded_files:
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            text = find_text(filepath)
            file_details.append([filename, text])

    return render_template('display.html', files=file_details)


@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'],
                               filename)

@app.route('/download', methods=['POST'])
def save_file():
    option = request.form['action']
    texts = request.form.getlist("message[]")

    if option == "word":
      document = Document()
      for text in texts:
        document.add_paragraph(text.replace("\n", ""))
      document.save('extracted_text.docx')
      return send_file("extracted_text.docx",as_attachment=True, cache_timeout=0)

    elif option == "pdf":
      texts = "\n".join(texts)
      pdf = FPDF()
      pdf.alias_nb_pages()
      pdf.add_page()
      pdf.set_font('Times', '', 10)
      for text in texts.split("\n"):
        pdf.cell(0, 8, text.encode('utf-8').decode('latin-1'), 0, 1)
      pdf.output('extracted_text.pdf', 'F')
      return send_file("extracted_text.pdf",as_attachment=True, cache_timeout=0)

    elif option == "txt":
      with open("extracted_text.txt", "w") as f:
        for text in texts:
          f.write(text.replace("\n", ""))
      return send_file("extracted_text.txt",as_attachment=True, cache_timeout=0)

    return ""


if __name__ == "__main__":
	app.run(debug=True)